import json
import shutil
import subprocess
import sys
import textwrap

import pytest

from sample_factory import TABLE1, TABLE2_BODY, find_font, make_sample
from tableocr.cli import main
from tableocr.export import to_delimited, write_xlsx
from tableocr.image_io import load_pages
from tableocr.ocr import OcrOptions, clean_text, ocr_table, read_title, title_matches
from tableocr.table_detect import deskew, detect_tables

needs_font = pytest.mark.skipif(find_font() is None, reason="日本語フォントがありません")
needs_tesseract = pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract がありません")

EXPECTED_TABLE2 = [
    ["支店", "売上", "", "備考"],
    ["", "4月", "5月", ""],
    *TABLE2_BODY,
]


@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    path = tmp_path_factory.mktemp("img") / "sample.tif"
    make_sample(str(path))
    return str(path)


@pytest.fixture(scope="module")
def page_tables(sample):
    page = load_pages(sample)[0]
    gray, angle = deskew(page.gray)
    tables, masks = detect_tables(gray, page.dpi, page=page.index)
    return page, gray, angle, tables, masks


@needs_font
def test_multipage_load(sample):
    pages = load_pages(sample)
    assert len(pages) == 2
    assert pages[0].dpi == 300


@needs_font
def test_deskew_detects_rotation(page_tables):
    _, _, angle, _, _ = page_tables
    assert angle == pytest.approx(-0.6, abs=0.15)


@needs_font
def test_detects_both_tables(page_tables):
    _, _, _, tables, _ = page_tables
    assert [(t.n_rows, t.n_cols) for t in tables] == [(3, 3), (5, 4)]


@needs_font
def test_merged_cells(page_tables):
    _, _, _, tables, _ = page_tables
    spans = {(c.row, c.col): (c.row_span, c.col_span) for c in tables[1].cells}
    assert spans[(0, 0)] == (2, 1)  # 支店
    assert spans[(0, 1)] == (1, 2)  # 売上
    assert spans[(0, 3)] == (2, 1)  # 備考
    assert spans[(1, 1)] == (1, 1)
    assert len(tables[1].cells) == 4 * 5 - 3


@needs_font
def test_roi_limits_search(sample):
    page = load_pages(sample)[0]
    tables, _ = detect_tables(page.gray, page.dpi, roi=(100, 650, 1500, 700))
    assert len(tables) == 1
    assert tables[0].n_cols == 4


@needs_font
@needs_tesseract
def test_ocr_tables(page_tables):
    page, gray, _, tables, masks = page_tables
    options = OcrOptions(jobs=4)
    assert ocr_table(gray, masks, tables[0], page.dpi, options) == TABLE1
    assert ocr_table(gray, masks, tables[1], page.dpi, options) == EXPECTED_TABLE2


@needs_font
@needs_tesseract
def test_read_titles(page_tables):
    page, gray, _, tables, masks = page_tables
    titles = [read_title(gray, masks, t, page.dpi, OcrOptions()) for t in tables]
    # 本文中の「構成表」ではなく、表のすぐ上の表題を読むこと
    assert titles == ["(別表)", "(構成表)"]


def test_title_matches():
    assert title_matches("（構成表）", "構成表")
    assert title_matches("(構成表)", "（構成表）")
    assert title_matches("( 構 成 表 )", "構成表")
    assert not title_matches("(別表)", "構成表")
    assert not title_matches("", "構成表")


@needs_font
@needs_tesseract
def test_cli_default_picks_title_table(sample, tmp_path, capsys):
    out = tmp_path / "out.xlsx"
    assert main([sample, "-o", str(out)]) == 0
    from openpyxl import load_workbook

    ws = load_workbook(out).active
    assert ws["A3"].value == "東京"
    assert {str(r) for r in ws.merged_cells.ranges} == {"A1:A2", "B1:C1", "D1:D2"}


@needs_font
@needs_tesseract
def test_cli_title_not_found(sample):
    assert main([sample, "--title", "存在しない表題"]) == 1


@needs_font
@needs_tesseract
def test_cli_table_number_overrides_title(sample, capsys):
    assert main([sample, "--table", "1"]) == 0
    assert capsys.readouterr().out.splitlines()[0] == "No\tItem\tQty"


@needs_font
@needs_tesseract
def test_cli_keyword(sample, capsys):
    assert main([sample, "--keyword", "名古屋"]) == 0
    assert "名古屋" in capsys.readouterr().out


@needs_font
def test_cli_list(sample, capsys):
    assert main([sample, "--list"]) == 0
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 2
    assert "5 行 x 4 列" in lines[1]
    if shutil.which("tesseract"):
        assert "対象" in lines[1] and "対象" not in lines[0]


@needs_font
def test_cli_bad_table_number(sample):
    assert main([sample, "--table", "9"]) == 2


def test_clean_text():
    assert clean_text("名 古 屋\n") == "名古屋"
    assert clean_text("1, 200") == "1,200"
    assert clean_text("3 . 5 kg") == "3.5 kg"
    assert clean_text("Apple  Pie") == "Apple Pie"
    assert clean_text("一行目\n二行目") == "一行目 二行目"
    assert clean_text("一行目\n二行目", keep_newlines=True) == "一行目\n二行目"


def test_to_delimited_quotes_newlines():
    text = to_delimited([["a", "b\nc"], ["", "d"]])
    assert text == 'a\t"b\nc"\r\n\td\r\n'


def test_write_xlsx_keeps_text(tmp_path):
    path = tmp_path / "t.xlsx"
    write_xlsx(str(path), [["001", "x"]])
    from openpyxl import load_workbook

    assert load_workbook(path).active["A1"].value == "001"


def test_block_network_refuses_connections():
    code = textwrap.dedent(
        """
        import socket
        from tableocr.offline import block_network, NetworkBlockedError
        block_network()
        for call in (lambda: socket.create_connection(("127.0.0.1", 80)),
                     lambda: socket.getaddrinfo("example.com", 443),
                     lambda: socket.socket().connect(("127.0.0.1", 80))):
            try:
                call()
            except NetworkBlockedError:
                continue
            raise SystemExit("通信できてしまった")
        print("ok")
        """
    )
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "ok"


@needs_font
@needs_tesseract
def test_full_run_makes_no_network_access(sample, tmp_path):
    """実行中に起きた通信・外部プログラム起動を監査フックで全部記録し、通信が無いことを確かめる。"""
    code = textwrap.dedent(
        f"""
        import json, sys
        events = []
        def hook(name, args):
            if name.startswith("socket.") or name == "subprocess.Popen":
                events.append([name, str(args[0]) if args else ""])
        sys.addaudithook(hook)
        from tableocr.cli import main
        rc = main([{sample!r}, "-o", {str(tmp_path / "o.xlsx")!r}])
        print(json.dumps({{"rc": rc, "events": events}}))
        """
    )
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    report = json.loads(result.stdout.strip().splitlines()[-1])
    assert report["rc"] == 0
    assert not [e for e in report["events"] if e[0].startswith("socket.")]
    # 起動する外部プログラムは tesseract だけ
    launched = {e[1] for e in report["events"] if e[0] == "subprocess.Popen"}
    assert launched and all("tesseract" in p for p in launched)
