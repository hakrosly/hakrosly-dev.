"""読み取った表を TSV / CSV / Excel / クリップボードに出力する。"""

from __future__ import annotations

import csv
import io
import platform
import shutil
import subprocess

from .table_detect import Table


def to_delimited(grid: list[list[str]], delimiter: str = "\t") -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerows(grid)
    return buf.getvalue()


def write_text(path: str, grid: list[list[str]], delimiter: str) -> None:
    # Excel で文字化けしないよう BOM 付き UTF-8 で保存する
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        f.write(to_delimited(grid, delimiter))


def write_xlsx(path: str, grid: list[list[str]], table: Table | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "table"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r, row in enumerate(grid, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            # 先頭の 0 などが消えないよう文字列として入れる
            cell.number_format = "@"
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if table is not None:
        for cell in table.cells:
            if cell.row_span > 1 or cell.col_span > 1:
                ws.merge_cells(
                    start_row=cell.row + 1,
                    start_column=cell.col + 1,
                    end_row=cell.row + cell.row_span,
                    end_column=cell.col + cell.col_span,
                )
    wb.save(path)


def write_output(path: str, grid: list[list[str]], table: Table | None = None) -> None:
    lower = path.lower()
    if lower.endswith(".xlsx"):
        write_xlsx(path, grid, table)
    elif lower.endswith(".csv"):
        write_text(path, grid, ",")
    else:
        write_text(path, grid, "\t")


def copy_to_clipboard(text: str) -> None:
    """表をクリップボードに入れる。Excel で Ctrl+V するだけで貼れる。"""
    system = platform.system()
    if system == "Windows":
        # clip.exe は BOM 付き UTF-16 なら日本語も正しく受け取れる
        subprocess.run(["clip"], input=text.encode("utf-16"), check=True)
    elif system == "Darwin":
        subprocess.run(["pbcopy"], input=text.encode("utf-8"), check=True)
    elif shutil.which("xclip"):
        subprocess.run(["xclip", "-selection", "clipboard"], input=text.encode("utf-8"), check=True)
    elif shutil.which("wl-copy"):
        subprocess.run(["wl-copy"], input=text.encode("utf-8"), check=True)
    else:
        raise RuntimeError("クリップボードにコピーするコマンドが見つかりません (xclip か wl-copy を入れてください)")
